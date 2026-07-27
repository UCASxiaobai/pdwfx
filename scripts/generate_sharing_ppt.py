# -*- coding: utf-8 -*-
"""Generate pdwfx project sharing PowerPoint."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "docs" / "pdwfx-project-sharing.pptx"

BLUE = RGBColor(0x1D, 0x4E, 0xD8)
DARK = RGBColor(0x1F, 0x29, 0x37)
GRAY = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0x0E, 0x74, 0x90)


def set_slide_bg(slide, rgb=RGBColor(0xF8, 0xFA, 0xFC)):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb


def add_title_bar(slide, title, subtitle=None):
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(1.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()
    tf = bar.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = WHITE
    if subtitle:
        box = slide.shapes.add_textbox(Inches(0.55), Inches(1.2), Inches(9), Inches(0.45))
        p2 = box.text_frame.paragraphs[0]
        p2.text = subtitle
        p2.font.size = Pt(14)
        p2.font.color.rgb = GRAY


def add_bullets(slide, items, left=0.65, top=1.75, width=8.8, height=5.2, font_size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            text, level = item
            p.text = text
            p.level = level
        else:
            p.text = item
            p.level = 0
        p.font.size = Pt(font_size if p.level == 0 else font_size - 2)
        p.font.color.rgb = DARK
        p.space_after = Pt(8)


def add_two_col(slide, left_title, left_items, right_title, right_items):
    for title, items, x in [
        (left_title, left_items, 0.55),
        (right_title, right_items, 5.15),
    ]:
        h = slide.shapes.add_textbox(Inches(x), Inches(1.55), Inches(4.2), Inches(0.4))
        p = h.text_frame.paragraphs[0]
        p.text = title
        p.font.bold = True
        p.font.size = Pt(16)
        p.font.color.rgb = ACCENT
        add_bullets(slide, items, left=x, top=1.95, width=4.2, height=4.8, font_size=15)


def title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BLUE)
    t = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(8.5), Inches(1.2))
    p = t.text_frame.paragraphs[0]
    p.text = "通信侦获信号分析系统"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    s = slide.shapes.add_textbox(Inches(0.8), Inches(3.2), Inches(8.5), Inches(0.8))
    p2 = s.text_frame.paragraphs[0]
    p2.text = "pdwfx 项目开发交流分享"
    p2.font.size = Pt(22)
    p2.font.color.rgb = RGBColor(0xBF, 0xDB, 0xFE)
    m = slide.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(8), Inches(1))
    p3 = m.text_frame.paragraphs[0]
    p3.text = "原始 CSV → 场景筛选 → 信号分析 → 可视化研判"
    p3.font.size = Pt(16)
    p3.font.color.rgb = WHITE


def content_slide(prs, title, bullets, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_title_bar(slide, title, subtitle)
    add_bullets(slide, bullets)


def build():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    title_slide(prs)

    content_slide(
        prs,
        "分享议程",
        [
            "1. 业务背景与建设目标",
            "2. 系统架构与技术栈",
            "3. 端到端业务流程（演示路径）",
            "4. 阶段一：方位场景发现",
            "5. 阶段二：信号分析研判",
            "6. 阶段三：可视化与跨频关联",
            "7. 工程化设计与踩坑案例",
            "8. 接口集成与后续方向",
        ],
    )

    content_slide(
        prs,
        "业务背景",
        [
            "侦察平台持续产出海量 PDW/CSV，含多频点、多目标、轮询/同频/异频等复杂模式",
            "人工在全量数据中找「值得分析」的时段成本高、易遗漏",
            "建设目标：自动化筛场景 + 分网编批 + 平台/主从/通信链研判 + 图表支撑",
            "核心思路：先缩小样本（优质场景），再深度分析（信号分析）",
        ],
        subtitle="Why pdwfx?",
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_title_bar(slide, "三段式系统定位", "前置 · 核心 · 后置")
    add_two_col(
        slide,
        "前置：方位场景发现",
        [
            "从 CSV 找优质时段",
            "连续轨迹场景",
            "多设备轮询场景",
            "输出 Top-K 场景表",
        ],
        "核心：信号分析",
        [
            "场景 CSV 分通信网",
            "目标编批与拆分",
            "平台类型 GROUND/AIR/AWACS",
            "主从 + 通信链 D01–D06",
        ],
    )
    box = slide.shapes.add_textbox(Inches(0.55), Inches(5.9), Inches(9), Inches(0.8))
    p = box.text_frame.paragraphs[0]
    p.text = "后置：Vue 可视化 + 汇总表 + Word 导出 + 跨频方位关联（分析后推断）"
    p.font.size = Pt(15)
    p.font.color.rgb = DARK

    content_slide(
        prs,
        "技术栈与仓库结构",
        [
            "后端：Java 8 + Spring Boot 2.7（单进程：com.scenefinder + com.pdwfx.signal）",
            "前端：Vue 3 + Vite + ECharts + OpenLayers",
            "部署：后端 :18080，前端 dev :5173，/api 代理",
            "pdwfx/",
            ("  backend/   — 场景发现 + 信号分析", 1),
            ("  frontend/  — SceneWorkflow 主工作流", 1),
            ("  docs/      — interface-api、algorithm-design", 1),
            ("  output/    — 场景汇总、导出 CSV、可视化 JSON", 1),
        ],
    )

    content_slide(
        prs,
        "端到端用户流程",
        [
            "① 上传数据：多 CSV 一次上传（PrcFf 等原始侦获表）",
            "② 场景筛选：配置频段、时间窗、TOP-K、频率容差 → analyze-upload",
            "③ 信号分析：逐场景导出 CSV → 进程内 Facade 分析 → 汇总表",
            "④ 结果查看：筛选表 / 方位图 / 网络研判 / 跨频关联 / Word 导出",
            "一键入口：「开始分析（场景筛选 → 信号分析）」",
        ],
        subtitle="SceneWorkflow 四步",
    )

    content_slide(
        prs,
        "关键 CSV 字段语义",
        [
            "pl        — 频率 MHz（建轨、频段过滤）",
            "xhfw      — 目标方位角 °（轨迹聚类；全是目标数据）",
            "zcsj      — 侦测时间（分帧、时间窗）",
            "zjwzjd/wd — 本机（侦察平台）经纬度",
            "nSignalTime — 驻留（导出保留，供占空比/通信链）",
            "注意：方位轨迹图标注「目标1、目标2…」，不将 xhfw 误标为本机",
        ],
    )

    content_slide(
        prs,
        "阶段一：建轨流程",
        [
            "读入检测点 → 按 freqClusterGapMhz 分频段（默认 0.01，与页面容差一致）",
            "每频段：frameSeconds(0.5s) 分帧 → 帧内 4° 方位聚类",
            "跨帧：预测方位 + 6° 门控 + 频率门控，贪心关联",
            "丢帧 > maxMissedFrames(2) 则断轨",
            "过滤：≥ minTrackPoints(20) 且 ≥ minTrackSeconds(30s)",
        ],
        subtitle="TrackBuilderService",
    )

    content_slide(
        prs,
        "阶段一：场景评分公式",
        [
            "同频块内滑动窗：120s 窗长，30s 步进",
            "score = 2.5×轨迹数 + 2.5×方位分离 + 1.5×密度 + 0.75×设备数 − 平滑惩罚",
            "若 medianSeparation < 3° 再 −0.35",
            "偏向：多设备轮询（8+）、同频多轨并行（~4+）",
            "单轨平缓场景得分偏低（~2.2），宽频段 TOP-K 下可能落选",
        ],
        subtitle="SceneScorerService",
    )

    content_slide(
        prs,
        "阶段一：融合与输出",
        [
            "连续轨迹 TOP-K 与轮询 TOP-K 分项配额（默认各 50）",
            "按 score 统一排序，时间重合场景融合去重（阈值 0.92）",
            "输出：scene_summary.csv、轨迹检测 CSV、visualization-data.json",
            "频率容差统一：页面 freqTolerance 覆盖建轨 + 场景分组 + 信号分网",
        ],
    )

    content_slide(
        prs,
        "阶段二：信号分析流水线",
        [
            "导入 PrcFf → DetectSignal（PL/XHFW/XHFD/ZCSJ 等映射）",
            "按 freqTolerance 分网 → SAME_FREQ / MULTI_FREQ / HOPPING",
            "单网：方位轨迹(DBSCAN+RANSAC+样条) → 幅度二次聚类",
            "运动分类 → 通信节奏(PRI/burst/占空比) → 平台类型",
            "主从研判 → 通信链 D01–D06 → NetworkView/TargetView",
        ],
        subtitle="SignalAnalysisService",
    )

    content_slide(
        prs,
        "阶段二：研判输出",
        [
            "targetType：GROUND（地面站）/ AWACS（预警机）/ AIR（飞机）",
            "role：MASTER / SLAVE（发射时序为主判据）",
            "图表序列：azimuthSeries、signalSeries、rawAzimuthSeries…",
            "通信链：波道规则 + networkType + 证据链文字",
            "逐场景 process-scene，避免大批量 preload 超时",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_title_bar(slide, "阶段三：三层方位可视化", "由浅入深")
    add_two_col(
        slide,
        "场景筛选方位图",
        ["SceneBearingViz", "建轨代表轨迹", "目标1、目标2…"],
        "信号分析方位图",
        ["SceneAnalysisBearingViz", "NetworkView 方位序列", "目标1-飞机 等标签"],
    )
    add_bullets(
        slide,
        [
            "跨频方位关联（SceneCrossFreqMatchViz）",
            ("  分析后推断，不改变编批结果", 1),
            ("  合成轨迹：线段颜色 = 该时段频率", 1),
            ("  分轨视图：颜色 = 推断目标，线型 = 频率", 1),
        ],
        top=5.2,
        font_size=15,
    )

    content_slide(
        prs,
        "单网研判工作台",
        [
            "AnalystWorkbench：时间-方位、幅度、频率、原始散点四宫格",
            "通信节奏区：TOA、PRI 分布、Jitter、占空比趋势",
            "OpenLayers 地图：方位射线、选中目标联动",
            "多图点击联动：选中目标加粗、同时间点钻石标记",
            "图例优化：Y 轴名称居中，避免与顶部图例重叠",
        ],
    )

    content_slide(
        prs,
        "跨频方位匹配参数",
        [
            "方位门限 (°)：同时刻两轨方位差上限",
            "时间对齐 (ms)：两轨采样时刻允许偏差",
            "最少匹配点：判定关联所需重合点数",
            "最短重合 (ms)：两轨时间重叠下限",
            "可选：运动趋势一致、显示未关联轨迹（灰色）",
            "时间重合场景自动合并（传递闭包）",
        ],
        subtitle="bearingMatch.js · 后置推断",
    )

    content_slide(
        prs,
        "工程化设计要点",
        [
            "in-process 部署：ScenePipeline 直接调 SignalAnalysisFacade",
            "会话 LRU(256) + ensure-session：驱逐后可从 CSV 恢复",
            "SceneSourceExportService：按行提取，保留原始表头全部列",
            "analysisViewCache：避免 deep watch 导致重复 abort/loading 卡死",
            "大可视化：超阈值按需拉取 + 散点降采样",
        ],
    )

    content_slide(
        prs,
        "踩坑案例（分享讨论）",
        [
            "场景 240.55 MHz，分析结果却是 367 MHz → 导出 CSV 路径错误",
            "宽频段找不到 240.55，窄频段可以 → TOP-K 竞争 + 单轨低分",
            "30 MHz 场景分组 vs 0.01 分网容差不一致 → 已统一 freqTolerance",
            "方位图误标「本机」→ xhfw 均为目标方位，已改目标1/2",
            "统计图图例与 Y 轴名重叠 → 布局与轴名位置优化",
        ],
    )

    content_slide(
        prs,
        "关键 API（集成参考）",
        [
            "POST /api/scenes/analyze-upload     — 上传场景筛选",
            "POST /api/scenes/process-scene      — 单场景导出+分析",
            "GET  /api/scenes/visualization-data — 方位可视化数据",
            "POST /api/signals/analyze           — 直接分析 CSV",
            "GET  /api/signals/analysis/{id}/networks/{networkId}",
            "详见 docs/interface-api.md、external-integration.md",
        ],
    )

    content_slide(
        prs,
        "默认参数速查",
        [
            "频段：200–1000 MHz | 评分窗 120s/步进 30s",
            "TOP-K 连续轨迹/轮询：各 50 | 最少轨迹数 1",
            "频率容差：0.01 MHz（建轨+场景+分网共用）",
            "建轨：分帧 0.5s | 方位聚类 4° | 关联门控 6°",
            "最短轨迹 30s / 20 点 | 场景分离软阈值 3°",
        ],
    )

    content_slide(
        prs,
        "Demo 建议路径",
        [
            "1. 上传 PrcFf1139.csv",
            "2. 频段 200–1000 MHz，频率容差 0.01，开始分析",
            "3. 查看汇总表：场景频点 vs 网络频率",
            "4. 打开场景方位图 → 信号分析分频方位图",
            "5. 点击行进入 AnalystWorkbench 多图联动",
            "6. 跨频关联：合成轨迹 + 调整匹配参数",
        ],
    )

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BLUE)
    t = slide.shapes.add_textbox(Inches(2), Inches(2.8), Inches(6), Inches(1.2))
    p = t.text_frame.paragraphs[0]
    p.text = "Q & A"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = WHITE
    s = slide.shapes.add_textbox(Inches(1.5), Inches(4.2), Inches(7), Inches(0.6))
    p2 = s.text_frame.paragraphs[0]
    p2.text = "感谢聆听"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(0xBF, 0xDB, 0xFE)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    build()
